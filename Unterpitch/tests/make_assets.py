"""Generates manual-test assets for every case in cases.json.

  decks/<id>.pptx    the slides (load into the app / extension)
  scripts/<id>.txt   what to say, with [Folie n] / [Slide n] markers (read it aloud while rehearsing)
  audio/<id>.wav     the script spoken by a Windows voice (only if a voice for that language is installed)

Usage: python tests/make_assets.py      (needs: pip install python-pptx; audio needs Windows)
"""
import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).parent
cases = json.loads((ROOT / "cases.json").read_text(encoding="utf-8"))
for d in ("decks", "scripts", "audio"):
    (ROOT / d).mkdir(exist_ok=True)


def make_deck(c, path):
    prs = Presentation()
    for s in c["slides"]:
        layout = prs.slide_layouts[1] if s["points"] else prs.slide_layouts[5]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s["title"]
        if s["points"]:
            tf = slide.placeholders[1].text_frame
            tf.text = s["points"][0]
            for p in s["points"][1:]:
                tf.add_paragraph().text = p
    prs.save(path)


def make_script(c):
    label = "Folie" if c["lang"] == "de" else "Slide"
    lines, last = [f"# {c['id']}: {c['about']}", ""], None
    for seg in c["segments"]:
        if seg["slide"] != last:
            lines += ["", f"[{label} {seg['slide']}]"]
            last = seg["slide"]
        lines.append(seg["text"])
    if not c["segments"]:
        lines.append("(stay silent for ~10 seconds)")
    return "\n".join(lines).strip() + "\n"


PS_SPEAK = r"""
Add-Type -AssemblyName System.Speech
$s = New-Object System.Speech.Synthesis.SpeechSynthesizer
$v = $s.GetInstalledVoices() | Where-Object { $_.VoiceInfo.Culture.Name -like '{culture}*' } | Select-Object -First 1
if (-not $v) { exit 3 }
$s.SelectVoice($v.VoiceInfo.Name)
$s.Rate = {rate}
$s.SetOutputToWaveFile('{out}')
$s.Speak([IO.File]::ReadAllText('{txt}', [Text.Encoding]::UTF8))
$s.Dispose()
"""


def make_audio(c, out):
    if not c["segments"] or sys.platform != "win32":
        return "skipped"
    txt = ROOT / "audio" / f"{c['id']}.speak.txt"
    txt.write_text(" ".join(s["text"] for s in c["segments"]), encoding="utf-8")
    # SAPI rate: -10..10; map the case's words-per-minute roughly onto it
    rate = max(-10, min(10, round((c["wpm"] - 150) / 12)))
    culture = "de" if c["lang"] == "de" else "en"
    ps = PS_SPEAK.replace("{culture}", culture).replace("{rate}", str(rate)).replace("{out}", str(out)).replace("{txt}", str(txt))
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True)
    txt.unlink(missing_ok=True)
    return "ok" if r.returncode == 0 else ("no voice" if r.returncode == 3 else "error")


summary = {}
for c in cases:
    make_deck(c, ROOT / "decks" / f"{c['id']}.pptx")
    (ROOT / "scripts" / f"{c['id']}.txt").write_text(make_script(c), encoding="utf-8")
    summary[c["id"]] = make_audio(c, ROOT / "audio" / f"{c['id']}.wav")

for k, v in summary.items():
    print(f"{k:24} deck ✓  script ✓  audio: {v}")
if "no voice" in summary.values():
    print("\nNo German Windows voice installed → German audio skipped. Add one under Settings → Time & language → Speech → Add voices (Deutsch), then rerun.")
